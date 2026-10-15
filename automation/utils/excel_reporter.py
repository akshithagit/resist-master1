"""
ResistAI - Enterprise Excel Test Report Generator
Uses openpyxl to generate comprehensive multi-sheet Excel workbooks:
1. Automation_Test_Report.xlsx (6 sheets: Executed, Passed, Failed, Skipped, Metrics, Defect Summary)
2. Passed_Test_Cases.xlsx
3. Failed_Test_Cases.xlsx
4. Summary_Report.xlsx
"""

from typing import List, Dict, Any
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from automation.config.config import (
    AUTOMATION_TEST_REPORT_XLSX,
    PASSED_TEST_CASES_XLSX,
    FAILED_TEST_CASES_XLSX,
    SUMMARY_REPORT_XLSX,
)
from automation.utils.logger import logger

# Professional Color Palette
NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Segoe UI", size=10)
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1B365D")

PASS_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
PASS_FONT = Font(name="Segoe UI", size=10, bold=True, color="155724")

FAIL_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
FAIL_FONT = Font(name="Segoe UI", size=10, bold=True, color="721C24")

SKIP_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
SKIP_FONT = Font(name="Segoe UI", size=10, bold=True, color="856404")

THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

class ExcelReportGenerator:
    @staticmethod
    def generate_all_reports(test_results: List[Dict[str, Any]], summary_stats: Dict[str, Any]):
        """
        Generates Automation_Test_Report.xlsx and individual split reports.
        """
        logger.info("Generating Excel automation test reports...")
        ExcelReportGenerator._generate_master_workbook(test_results, summary_stats)
        ExcelReportGenerator._generate_passed_workbook(test_results)
        ExcelReportGenerator._generate_failed_workbook(test_results)
        ExcelReportGenerator._generate_summary_workbook(summary_stats)
        logger.info("All Excel reports successfully generated.")

    @staticmethod
    def _generate_master_workbook(results: List[Dict[str, Any]], stats: Dict[str, Any]):
        wb = openpyxl.Workbook()
        
        # Sheet 1: Executed Test Cases
        ws1 = wb.active
        ws1.title = "Executed Test Cases"
        headers1 = ["Test ID", "Module", "Test Name", "Status", "Execution Time (s)", "Priority"]
        ws1.append(headers1)
        for r in results:
            ws1.append([
                r.get("test_id", ""),
                r.get("module", ""),
                r.get("title", ""),
                r.get("status", "PASSED"),
                round(r.get("duration", 0.0), 3),
                r.get("priority", "P1")
            ])
        ExcelReportGenerator._style_table(ws1, headers1)

        # Sheet 2: Passed Tests
        ws2 = wb.create_sheet(title="Passed Tests")
        headers2 = ["Test ID", "Module", "Test Name", "Duration (s)", "Priority"]
        ws2.append(headers2)
        for r in results:
            if r.get("status") == "PASSED":
                ws2.append([r.get("test_id"), r.get("module"), r.get("title"), round(r.get("duration", 0.0), 3), r.get("priority")])
        ExcelReportGenerator._style_table(ws2, headers2, default_fill=PASS_FILL)

        # Sheet 3: Failed Tests
        ws3 = wb.create_sheet(title="Failed Tests")
        headers3 = ["Test ID", "Module", "Test Name", "Failure Reason", "Screenshot Path"]
        ws3.append(headers3)
        for r in results:
            if r.get("status") == "FAILED":
                ws3.append([r.get("test_id"), r.get("module"), r.get("title"), r.get("error_message", "N/A"), r.get("screenshot_path", "")])
        ExcelReportGenerator._style_table(ws3, headers3, default_fill=FAIL_FILL)

        # Sheet 4: Skipped Tests
        ws4 = wb.create_sheet(title="Skipped Tests")
        headers4 = ["Test ID", "Module", "Test Name", "Reason"]
        ws4.append(headers4)
        for r in results:
            if r.get("status") == "SKIPPED":
                ws4.append([r.get("test_id"), r.get("module"), r.get("title"), r.get("skip_reason", "Not Applicable")])
        ExcelReportGenerator._style_table(ws4, headers4, default_fill=SKIP_FILL)

        # Sheet 5: Execution Metrics
        ws5 = wb.create_sheet(title="Execution Metrics")
        ws5.append(["ResistAI Automated Test Execution Metrics"])
        ws5.append([])
        metrics = [
            ("Total Test Cases", stats.get("total", len(results))),
            ("Passed Tests", stats.get("passed", 0)),
            ("Failed Tests", stats.get("failed", 0)),
            ("Skipped Tests", stats.get("skipped", 0)),
            ("Pass Rate (%)", f"{stats.get('pass_rate', 0.0):.2f}%"),
            ("Total Execution Time", f"{stats.get('total_duration', 0.0):.2f} s"),
            ("Average Test Duration", f"{stats.get('avg_duration', 0.0):.3f} s"),
            ("Target Deployment URL", stats.get("base_url", "N/A")),
            ("Execution Timestamp", stats.get("timestamp", "N/A")),
        ]
        for label, val in metrics:
            ws5.append([label, val])
        ws5.column_dimensions["A"].width = 28
        ws5.column_dimensions["B"].width = 40
        ws5["A1"].font = TITLE_FONT

        # Sheet 6: Defect Summary
        ws6 = wb.create_sheet(title="Defect Summary")
        headers6 = ["Defect ID", "Associated Test ID", "Module", "Severity", "Description", "Stack Trace Summary"]
        ws6.append(headers6)
        defect_idx = 1
        for r in results:
            if r.get("status") == "FAILED":
                ws6.append([
                    f"DEF-{defect_idx:03d}",
                    r.get("test_id"),
                    r.get("module"),
                    "High" if r.get("priority") == "P1" else "Medium",
                    r.get("error_message", "Unknown Assertion Failure"),
                    r.get("stack_trace", "N/A")[:200]
                ])
                defect_idx += 1
        ExcelReportGenerator._style_table(ws6, headers6)

        wb.save(str(AUTOMATION_TEST_REPORT_XLSX))

    @staticmethod
    def _generate_passed_workbook(results: List[Dict[str, Any]]):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Passed Test Cases"
        headers = ["Test ID", "Module", "Test Name", "Execution Time (s)", "Priority"]
        ws.append(headers)
        for r in results:
            if r.get("status") == "PASSED":
                ws.append([r.get("test_id"), r.get("module"), r.get("title"), round(r.get("duration", 0.0), 3), r.get("priority")])
        ExcelReportGenerator._style_table(ws, headers)
        wb.save(str(PASSED_TEST_CASES_XLSX))

    @staticmethod
    def _generate_failed_workbook(results: List[Dict[str, Any]]):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Failed Test Cases"
        headers = ["Test ID", "Module", "Test Name", "Error Message", "Screenshot"]
        ws.append(headers)
        for r in results:
            if r.get("status") == "FAILED":
                ws.append([r.get("test_id"), r.get("module"), r.get("title"), r.get("error_message", ""), r.get("screenshot_path", "")])
        ExcelReportGenerator._style_table(ws, headers, default_fill=FAIL_FILL)
        wb.save(str(FAILED_TEST_CASES_XLSX))

    @staticmethod
    def _generate_summary_workbook(stats: Dict[str, Any]):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Execution Summary"
        headers = ["Metric", "Value"]
        ws.append(headers)
        for k, v in stats.items():
            ws.append([str(k).replace("_", " ").title(), str(v)])
        ExcelReportGenerator._style_table(ws, headers)
        wb.save(str(SUMMARY_REPORT_XLSX))

    @staticmethod
    def _style_table(ws, headers: List[str], default_fill = None):
        # Header Styling
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = NAVY_HEADER
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Body Styling & Column Width auto-fit
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if default_fill:
                    cell.fill = default_fill

                # Special status highlighting
                val_str = str(cell.value)
                if val_str == "PASSED":
                    cell.fill = PASS_FILL
                    cell.font = PASS_FONT
                elif val_str == "FAILED":
                    cell.fill = FAIL_FILL
                    cell.font = FAIL_FONT
                elif val_str == "SKIPPED":
                    cell.fill = SKIP_FILL
                    cell.font = SKIP_FONT

        # Adjust Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
